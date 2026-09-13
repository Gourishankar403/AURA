import re
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
import os

def convert():
    doc = Document()
    
    # Read the markdown file
    md_path = r"C:\Users\gouri\.gemini\antigravity\brain\8a903ef1-49bb-44dd-b713-5adabb742c99\AURA_Comprehensive_Documentation.md"
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()
        
    in_code_block = False
    in_table = False
    table_data = []

    for line in lines:
        line = line.strip()
        if not line:
            continue
            
        # Handle code blocks (like the Mermaid diagram)
        if line.startswith('`'):
            in_code_block = not in_code_block
            if in_code_block:
                doc.add_paragraph("--- Diagram / Code Block ---").bold = True
            continue
            
        if in_code_block:
            p = doc.add_paragraph(line)
            try:
                p.runs[0].font.name = 'Courier New'
            except:
                pass
            continue
            
        # Handle Tables
        if line.startswith('|'):
            in_table = True
            if "---" not in line:
                cols = [c.strip() for c in line.split('|') if c.strip()]
                table_data.append(cols)
            continue
        else:
            if in_table and table_data:
                # Render the stored table
                table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
                table.style = 'Table Grid'
                for i, row in enumerate(table_data):
                    for j, cell in enumerate(row):
                        if j < len(table.columns):
                            table.cell(i, j).text = cell.replace('**', '')
                table_data = []
                in_table = False
            
        # Handle Headings
        if line.startswith('# '):
            doc.add_heading(line[2:].replace('**', ''), level=1)
        elif line.startswith('## '):
            doc.add_heading(line[3:].replace('**', ''), level=2)
        elif line.startswith('### '):
            doc.add_heading(line[4:].replace('**', ''), level=3)
        # Handle Lists
        elif line.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            parts = re.split(r'(\*\*.*?\*\*)', line[2:])
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    p.add_run(part[2:-2]).bold = True
                else:
                    p.add_run(part)
        # Handle Normal Paragraphs
        else:
            if line == "---":
                continue # Skip markdown horizontal rules
            p = doc.add_paragraph()
            parts = re.split(r'(\*\*.*?\*\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    p.add_run(part[2:-2]).bold = True
                else:
                    p.add_run(part)

    # Catch dangling table at the end of file
    if in_table and table_data:
        table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
        table.style = 'Table Grid'
        for i, row in enumerate(table_data):
            for j, cell in enumerate(row):
                if j < len(table.columns):
                    table.cell(i, j).text = cell.replace('**', '')

    out_path = r"C:\Users\gouri\Downloads\]RAG_POST\AURA_Project\AURA_Comprehensive_Documentation.docx"
    doc.save(out_path)
    print(f"Successfully generated Word document at {out_path}")

if __name__ == '__main__':
    convert()
